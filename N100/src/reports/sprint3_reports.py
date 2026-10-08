"""Sprint 3 report generator: screener Excel, peer Excel and radar charts."""
from __future__ import annotations
from pathlib import Path
import sqlite3
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from src.screener.engine import run, load_config, SCREEN_YEAR
from src.analytics.peer import compute_peer_percentiles, METRICS

ROOT=Path(__file__).resolve().parents[2]
DB=ROOT/"data"/"nifty100.db"
OUT=ROOT/"output"; REPORTS=ROOT/"reports"/"radar_charts"
GREEN=PatternFill("solid",fgColor="C6EFCE"); RED=PatternFill("solid",fgColor="FFC7CE"); YELLOW=PatternFill("solid",fgColor="FFEB9C"); GOLD=PatternFill("solid",fgColor="FFD966")

KPI_COLUMNS=[
("ROE %","return_on_equity_pct"),("ROCE %","roce_pct"),("NPM %","net_profit_margin_pct"),("OPM %","operating_profit_margin_pct"),
("D/E","debt_to_equity"),("FCF Cr","free_cash_flow_cr"),("FCF CAGR 5Y %","fcf_cagr_5yr"),("CFO/PAT","cfo_pat_ratio"),
("FCF Positive","fcf_positive_flag"),("Revenue CAGR 5Y %","revenue_cagr_5yr"),("PAT CAGR 5Y %","pat_cagr_5yr"),("EPS CAGR 5Y %","eps_cagr_5yr"),
("ICR","interest_coverage"),("Asset Turnover","asset_turnover"),("P/E","pe_ratio"),("P/B","pb_ratio"),
("Dividend Yield %","dividend_yield_pct"),("Market Cap Cr","market_cap_crore"),("Net Profit Cr","net_profit"),("Sales Cr","sales")]

def _save_screener_excel(results):
    OUT.mkdir(exist_ok=True); path=OUT/"screener_output.xlsx"
    with pd.ExcelWriter(path,engine="openpyxl") as writer:
        for name,df in results.items():
            cols=["company_id","company_name","broad_sector","strict_match"]+[c for _,c in KPI_COLUMNS]+["composite_quality_score"]
            x=df[[c for c in cols if c in df.columns]].copy()
            x.columns=[{"company_id":"Company ID","company_name":"Company Name","broad_sector":"Sector","strict_match":"Strict Match", "composite_quality_score":"Composite Score"}.get(c,dict(KPI_COLUMNS).get(c,c)) for c in x.columns]
            x.to_excel(writer,sheet_name=name[:31],index=False)
    cfg=load_config()
    aliases={"roe_min":"ROE %","de_max":"D/E","fcf_min":"FCF Cr","revenue_cagr_5yr_min":"Revenue CAGR 5Y %","pat_cagr_5yr_min":"PAT CAGR 5Y %","opm_min":"OPM %","pe_max":"P/E","pb_max":"P/B","dividend_yield_min":"Dividend Yield %","icr_min":"ICR","market_cap_min":"Market Cap Cr","net_profit_min":"Net Profit Cr","eps_cagr_min":"EPS CAGR 5Y %","asset_turnover_min":"Asset Turnover","sales_min":"Sales Cr"}
    wb=load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=Font(bold=True); cell.alignment=Alignment(horizontal="center")
        headers={c.value:c.column for c in ws[1]}
        filters=cfg["presets"].get(ws.title,{ }).get("filters",{})
        for row in ws.iter_rows(min_row=2):
            strict=row[headers.get("Strict Match",1)-1].value if "Strict Match" in headers else True
            sector=row[headers["Sector"]-1].value if "Sector" in headers else ""
            for key,threshold in filters.items():
                colname=aliases.get(key);
                if not colname or colname not in headers: continue
                cell=row[headers[colname]-1]; val=cell.value
                passes=False
                if key=="de_max" and sector=="Financials": passes=True
                elif key=="icr_min" and (val is None or (isinstance(val,float) and np.isnan(val))): passes=True
                elif isinstance(val,(int,float)) and val is not None:
                    passes = val > threshold if key.endswith("_min") else val < threshold
                cell.fill=GREEN if passes else RED
            if strict is False and "Strict Match" in headers: row[headers["Strict Match"]-1].fill=YELLOW
        for col in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=min(28,max(12,max(len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,30)+1))+2))
    wb.save(path); return path

def _peer_excel():
    con=sqlite3.connect(DB)
    pct=pd.read_sql_query("SELECT * FROM peer_percentiles WHERE year=?",con,params=[SCREEN_YEAR])
    peers=pd.read_sql_query("SELECT * FROM peer_groups",con)
    names=pd.read_sql_query("SELECT id company_id,company_name FROM companies",con)
    ratios=pd.read_sql_query("SELECT company_id,year,composite_quality_score FROM financial_ratios WHERE year=?",con,params=[SCREEN_YEAR])
    con.close()
    pct=pct.merge(peers,on=["company_id","peer_group_name"]).merge(names,on="company_id").merge(ratios,on=["company_id","year"],how="left")
    path=OUT/"peer_comparison.xlsx"
    with pd.ExcelWriter(path,engine="openpyxl") as writer:
        for group in sorted(pct.peer_group_name.unique()):
            g=pct[pct.peer_group_name==group]
            values=g.pivot_table(index=["company_id","company_name","is_benchmark"],columns="metric",values="value",aggfunc="first").reset_index()
            ranks=g.pivot_table(index=["company_id"],columns="metric",values="percentile_rank",aggfunc="first").reset_index()
            values.columns=[str(c) for c in values.columns]
            ranks.columns=[str(c)+" Percentile" if c!="company_id" else c for c in ranks.columns]
            final=values.merge(ranks,on="company_id",how="left")
            final.to_excel(writer,sheet_name=group[:31],index=False)
    wb=load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        headers={c.value:c.column for c in ws[1]}
        for c in ws[1]: c.font=Font(bold=True); c.alignment=Alignment(horizontal="center")
        for row in range(2,ws.max_row+1):
            bench=ws.cell(row,headers.get("is_benchmark",1)).value
            if bench==1:
                for c in range(1,ws.max_column+1): ws.cell(row,c).fill=GOLD
            for h,col in headers.items():
                if isinstance(h,str) and h.endswith(" Percentile"):
                    v=ws.cell(row,col).value
                    if isinstance(v,(int,float)):
                        ws.cell(row,col).fill=GREEN if v>=.75 else RED if v<=.25 else YELLOW
        # summary row: median of numeric metric/value and percentile columns
        sr=ws.max_row+2; ws.cell(sr,1,"Peer Group Median")
        for col in range(1,ws.max_column+1):
            vals=[ws.cell(r,col).value for r in range(2,ws.max_row-1) if isinstance(ws.cell(r,col).value,(int,float))]
            if vals: ws.cell(sr,col,float(np.median(vals)))
        for col in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(col)].width=min(26,max(12,max(len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,25)+1))+2))
    wb.save(path); return path

def _radar_charts():
    REPORTS.mkdir(parents=True,exist_ok=True)
    con=sqlite3.connect(DB)
    r=pd.read_sql_query("SELECT * FROM financial_ratios",con)
    peers=pd.read_sql_query("SELECT peer_group_name,company_id FROM peer_groups",con)
    companies=pd.read_sql_query("SELECT id company_id,company_name FROM companies",con)
    con.close()
    r["yn"]=pd.to_numeric(r["year"].str[:4],errors="coerce")
    # Use the common screening year where available; otherwise latest available ratio row.
    base=r[r["year"]==SCREEN_YEAR].copy()
    missing=companies.loc[~companies.company_id.isin(base.company_id),"company_id"]
    if len(missing):
        fallback=r[r.company_id.isin(missing)].sort_values("yn").groupby("company_id").tail(1)
        base=pd.concat([base,fallback],ignore_index=True)
    r=base.merge(companies,on="company_id",how="right").merge(peers,on="company_id",how="left")
    axes=[("ROE","return_on_equity_pct"),("ROCE","roce_pct"),("NPM","net_profit_margin_pct"),("D/E","debt_to_equity"),("FCF","free_cash_flow_cr"),("PAT CAGR","pat_cagr_5yr"),("Revenue CAGR","revenue_cagr_5yr"),("Composite","composite_quality_score")]
    for _,row in r.iterrows():
        group=row.peer_group_name if pd.notna(row.peer_group_name) else "No peer group assigned"
        g=r[r.peer_group_name.eq(group)] if group!="No peer group assigned" else r
        scaled={}
        for label,col in axes:
            x=pd.to_numeric(g[col],errors="coerce"); lo=x.quantile(.1); hi=x.quantile(.9)
            s=((x.clip(lo,hi)-lo)/(hi-lo)*100) if pd.notna(lo) and pd.notna(hi) and hi!=lo else pd.Series(50.0,index=g.index)
            if label=="D/E": s=100-s
            scaled[label]=s.fillna(50)
        if group=="No peer group assigned":
            # Nifty 100 average reference for standalone charts.
            avg=np.array([float(pd.to_numeric(r[col],errors="coerce").mean()) if pd.to_numeric(r[col],errors="coerce").notna().any() else 50 for _,col in axes])
        else:
            avg=np.array([float(np.mean(scaled[a])) for a,_ in axes])
        vals=[]
        for a,_ in axes:
            vals.append(float(scaled[a].loc[row.name]) if row.name in scaled[a].index else 50.0)
        vals=np.array(vals)
        theta=np.linspace(0,2*np.pi,len(axes),endpoint=False).tolist(); theta+=theta[:1]
        avg=np.r_[avg,avg[0]]; vals=np.r_[vals,vals[0]]
        fig,ax=plt.subplots(figsize=(7,7),subplot_kw=dict(polar=True))
        ax.plot(theta,avg,linestyle="--",linewidth=1.5,label="Nifty 100 Average" if group=="No peer group assigned" else "Peer Average")
        ax.fill(theta,vals,alpha=.18); ax.plot(theta,vals,linewidth=2,label=str(row.company_id))
        ax.set_xticks(theta[:-1]); ax.set_xticklabels([a for a,_ in axes],fontsize=9); ax.set_ylim(0,100); ax.set_title(f"{row.company_id} — {group}",pad=20); ax.legend(loc="upper right",bbox_to_anchor=(1.25,1.1))
        fig.tight_layout(); fig.savefig(REPORTS/f"{row.company_id}_radar.png",dpi=160); plt.close(fig)
    return len(list(REPORTS.glob("*_radar.png")))

def main():
    df,results=run()
    peer_rows=compute_peer_percentiles()
    sp=_save_screener_excel(results)
    pp=_peer_excel()
    charts=_radar_charts()
    return {"screener_sheets":len(results),"peer_percentile_rows":peer_rows,"radar_charts":charts,"screener":str(sp),"peer_comparison":str(pp),"preset_counts":{k:len(v) for k,v in results.items()}}

if __name__=="__main__": print(main())
